"""
Internal Assessment (IA) Sheets Generator for 4 Team Members (Python-docx).
Generates Team_Members_IA.docx containing formal continuous evaluation sheets,
role matrices, sprint diary logs, and assessment rubrics for all 4 team members.
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_paragraph(doc, text, size=15, bold=True, color=RGBColor(16, 44, 87), align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return p

def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    NAVY = RGBColor(16, 44, 87)

    members_data = [
        {
            "name": "Harsh Zhadge (Lead)",
            "roll": "TE-COMP-01",
            "prn": "72200101A",
            "role": "Team Lead & System Architect",
            "module": "FastAPI Gateway, Inference Pipeline Orchestration, Provider Adapters & Caching",
            "responsibilities": [
                "Architected overall decoupled system design and project repository layout.",
                "Built asynchronous FastAPI gateway endpoints (/api/health, /api/generate, /api/analyze).",
                "Engineered inference timing orchestration using microsecond-precision time.perf_counter().",
                "Implemented SQLite exact-match prompt cache reducing duplicate query costs to ₹0.00.",
                "Constructed offline DemoProvider and optional external APIProvider adapter."
            ]
        },
        {
            "name": "Team Member 2",
            "roll": "TE-COMP-02",
            "prn": "72200102B",
            "role": "ML & Routing Policy Engineer",
            "module": "Request Complexity Analyzer, Task Categorization & Privacy Protection Shield",
            "responsibilities": [
                "Designed deterministic syntactic complexity scoring algorithm (0.0 to 1.0 range).",
                "Formulated keyword heuristics across Coding, Reasoning, Mathematics, and General domains.",
                "Implemented Privacy Shield for regex matching of Aadhaar, PAN, passwords, and medical data.",
                "Built explainable multi-objective routing utility algorithm with configurable weights.",
                "Developed automated unit test suites for analyzer and routing services."
            ]
        },
        {
            "name": "Team Member 3",
            "roll": "TE-COMP-03",
            "prn": "72200103C",
            "role": "Data Science & Evaluation Specialist",
            "module": "Benchmark Dataset Engineering, Statistical EDA & ML Router Baseline",
            "responsibilities": [
                "Synthesized 325-sample benchmark dataset covering diverse prompt domains and complexity levels.",
                "Conducted exploratory data analysis (EDA) and generated publication-grade figures.",
                "Trained and evaluated Random Forest baseline classifier achieving 86.59% test accuracy.",
                "Formulated cost and latency reduction mathematical models and comparative statistics.",
                "Co-authored academic documentation and viva evaluation demonstration scripts."
            ]
        },
        {
            "name": "Team Member 4",
            "roll": "TE-COMP-04",
            "prn": "72200104D",
            "role": "Full-Stack UI & Telemetry Engineer",
            "module": "Web Interfaces (Router, Dashboard, Comparison UI) & Database Telemetry",
            "responsibilities": [
                "Developed clean, responsive academic web interfaces (AI Router, Dashboard, Comparison).",
                "Integrated Chart.js dynamic telemetry visualizations for model distribution and latency curves.",
                "Engineered side-by-side multi-model comparison view highlighting recommended tiers.",
                "Designed relational SQLite tables (requests, models, cache) and telemetry logging functions.",
                "Conducted cross-browser usability testing and assisted in compiling final report."
            ]
        }
    ]

    for idx, mem in enumerate(members_data):
        # Header block
        add_header_paragraph(doc, "SAVITRIBAI PHULE PUNE UNIVERSITY", size=13, bold=True)
        add_header_paragraph(doc, "DEPARTMENT OF COMPUTER ENGINEERING", size=11, bold=True)
        add_header_paragraph(doc, "PROBLEM BASED LEARNING (PBL) — INDIVIDUAL ASSESSMENT (IA) SHEET", size=12, bold=True, color=NAVY)
        add_header_paragraph(doc, "ACADEMIC YEAR 2026-2027 | SEMESTER - VI", size=10, bold=False)

        # Meta table
        table = doc.add_table(rows=5, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        info = [
            ("Project Group ID:", "TE-55", "Academic Year:", "2026-2027"),
            ("Project Title:", "Efficient AI Infrastructure for TESS — Cost- and Performance-Aware AI Model Router", "", ""),
            ("Student Name:", mem["name"], "Roll Number:", mem["roll"]),
            ("PRN Number:", mem["prn"], "Class & Div:", "T.E. Computer (Div A)"),
            ("Assigned Role:", mem["role"], "Guide Name:", "Project Guide / Professor")
        ]
        for r_i, r_val in enumerate(info):
            for c_i, c_val in enumerate(r_val):
                cell = table.rows[r_i].cells[c_i]
                set_cell_margins(cell, 60, 60, 80, 80)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(c_val)
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                if c_i in (0, 2):
                    r.font.bold = True
                    set_cell_background(cell, "F1F5F9")

        # Merge title row
        table.rows[1].cells[1].merge(table.rows[1].cells[2]).merge(table.rows[1].cells[3])

        # Section 1: Modular Ownership
        add_header_paragraph(doc, "\n1. Assigned Module & Key Deliverables", size=11, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT)
        p_mod = doc.add_paragraph()
        p_mod.paragraph_format.space_after = Pt(4)
        r = p_mod.add_run(f"Module: {mem['module']}")
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.bold = True

        for resp in mem["responsibilities"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.2)
            r = p.add_run(f"• {resp}")
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)

        # Section 2: Sprint Activity Log
        add_header_paragraph(doc, "\n2. Weekly Activity & Sprint Progress Log", size=11, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT)
        diary_table = doc.add_table(rows=5, cols=4)
        diary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        d_heads = ["Sprint / Week", "Planned Tasks & Deliverables", "Status / Outcome", "Guide Sign"]
        for c_i, h in enumerate(d_heads):
            cell = diary_table.rows[0].cells[c_i]
            set_cell_background(cell, "E2E8F0")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(h)
            run.font.bold = True
            run.font.size = Pt(9)

        sprints = [
            ("Weeks 1 - 2", "Literature survey, problem definition & system architecture mapping", "Completed (SRS & Design document approved)", ""),
            ("Weeks 3 - 4", f"Module implementation: {mem['module'][:50]}...", "Completed (Core algorithms & datasets verified)", ""),
            ("Weeks 5 - 6", "Full-stack integration with FastAPI, SQLite & Frontend UI", "Completed (End-to-end routing functional)", ""),
            ("Weeks 7 - 8", "Automated testing, EDA benchmarking & final PBL report compilation", "Completed (All acceptance criteria satisfied)", "")
        ]
        for r_i, s_row in enumerate(sprints, start=1):
            for c_i, s_val in enumerate(s_row):
                cell = diary_table.rows[r_i].cells[c_i]
                set_cell_margins(cell, 50, 50, 70, 70)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i in (0, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(s_val)
                run.font.name = "Times New Roman"
                run.font.size = Pt(8.5)

        # Section 3: Evaluation Rubric
        add_header_paragraph(doc, "\n3. Continuous Internal Assessment (IA) Evaluation Rubric", size=11, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT)
        rubric_table = doc.add_table(rows=6, cols=4)
        rubric_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        rub_heads = ["Assessment Parameter", "Evaluation Criteria", "Max Marks", "Marks Awarded"]
        for c_i, h in enumerate(rub_heads):
            cell = rubric_table.rows[0].cells[c_i]
            set_cell_background(cell, "E2E8F0")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(h)
            run.font.bold = True
            run.font.size = Pt(9)

        rubric_rows = [
            ("Problem Formulation & Review", "Literature survey, problem identification & objective clarity", "05", ""),
            ("System Architecture & Design", "Design modularity, database schema & algorithm formulation", "05", ""),
            ("Individual Implementation Rigor", "Code quality, assigned module completion & deliverables", "10", ""),
            ("Testing, Results & Presentation", "Validation, EDA charts, viva demonstration & project report", "05", ""),
            ("TOTAL MARKS", "Overall Cumulative Internal Assessment", "25", "")
        ]
        for r_i, rub_val in enumerate(rubric_rows, start=1):
            for c_i, val in enumerate(rub_val):
                cell = rubric_table.rows[r_i].cells[c_i]
                set_cell_margins(cell, 50, 50, 70, 70)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i in (2, 3) else WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(val)
                run.font.name = "Times New Roman"
                run.font.size = Pt(8.5)
                if r_i == 5:
                    run.font.bold = True
                    set_cell_background(cell, "F1F5F9")

        # Signatures
        doc.add_paragraph().paragraph_format.space_before = Pt(12)
        sig_p = doc.add_paragraph()
        sig_p.paragraph_format.space_after = Pt(4)
        sig_p.add_run("______________________________                                             ______________________________\n").font.bold = True
        sig_p.add_run("Signature of Student                                                                Signature of Project Guide").font.bold = True
        sig_p.runs[1].font.size = Pt(9.5)

        if idx < len(members_data) - 1:
            doc.add_page_break()

    out_file = "Team_Members_IA.docx"
    doc.save(out_file)
    print(f"[SUCCESS] Generated formal IA sheets for all 4 team members: {out_file}")

if __name__ == "__main__":
    main()
